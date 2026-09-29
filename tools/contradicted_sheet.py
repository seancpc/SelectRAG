#!/usr/bin/env python3
"""把 judge 結果中被判 contradicted / unsupported 的主張攤成人工複核表。

    python3 tools/contradicted_sheet.py results/judge/vllm_awq.json results/judge/review_awq.xlsx \\
        --ref results/judge/vllm_bf16_v2.json

欄位:claim_id、問題、參考段落、被判主張、judge 判定、引用 chunk、同題整段回答、
      對照組(--ref)同題回答與其判定摘要、你的判定、備註。
有 openpyxl 就寫 xlsx;沒有就寫同名 .csv(UTF-8 BOM,Excel 可直接開)。
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from typing import Any

HEADER = [
    "claim_id",
    "問題",
    "參考段落",
    "被判主張",
    "judge 判定",
    "引用 chunk",
    "同題整段回答",
    "對照組同題回答",
    "對照組判定摘要",
    "你的判定(支持/不支持)",
    "備註",
]
BAD = {"contradicted", "unsupported"}


def load(path: Path) -> dict[str, dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    return {str(a["id"]): a for a in data["answers"]}


def verdict_summary(a: dict[str, Any] | None) -> str:
    if a is None:
        return "(無此題)"
    if a["abstained"]:
        return "拒答"
    vs = [c.get("verdict") or "parse_fail" for c in a["claims"]]
    return f"{len(vs)} 條:" + "、".join(vs)


def rows_of(target: dict[str, dict[str, Any]], ref: dict[str, dict[str, Any]]) -> list[list[Any]]:
    out: list[list[Any]] = []
    for aid, a in target.items():
        if a["abstained"]:
            continue
        full = "".join(c["claim"] for c in a["claims"])
        r = ref.get(aid)
        if r is None:
            r_full = ""
        elif r["abstained"]:
            r_full = "(拒答)"
        else:
            r_full = "".join(c["claim"] for c in r["claims"])
        for c in a["claims"]:
            if c.get("verdict") in BAD:
                out.append(
                    [
                        c["claim_id"],
                        a["question"],
                        a["context"],
                        c["claim"],
                        c["verdict"],
                        c.get("evidence_chunk"),
                        full,
                        r_full,
                        verdict_summary(r),
                        "",
                        "",
                    ]
                )
    return out


def write(rows: list[list[Any]], out: Path) -> Path:
    if out.exists():
        raise FileExistsError(f"拒絕覆蓋:{out}")
    try:
        from openpyxl import Workbook  # type: ignore[import-untyped]
        from openpyxl.styles import Alignment  # type: ignore[import-untyped]
    except ImportError:
        out = out.with_suffix(".csv")
        if out.exists():
            raise FileExistsError(f"拒絕覆蓋:{out}") from None
        with out.open("w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(HEADER)
            w.writerows(rows)
        return out
    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    widths = [14, 30, 60, 40, 12, 10, 40, 40, 24, 14, 24]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = wd
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(out)
    return out


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("judge", type=Path)
    ap.add_argument("out", type=Path)
    ap.add_argument("--ref", type=Path, default=None)
    args = ap.parse_args(argv)
    target = load(args.judge)
    ref = load(args.ref) if args.ref else {}
    rows = rows_of(target, ref)
    path = write(rows, args.out)
    print(f"rows={len(rows)} 輸出:{path}")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main(sys.argv[1:]))
